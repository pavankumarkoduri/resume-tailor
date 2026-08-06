"""
resume_parser.py
================
Extract raw text from resume / job-description files.

Supported formats: .pdf and .docx only (per project scope).
"""
from __future__ import annotations
import io
import os


def extract_text(file_bytes: bytes, filename: str) -> str:
    """Extract plain text from a PDF or DOCX file given as raw bytes."""
    ext = os.path.splitext(filename)[1].lower()
    if ext == ".pdf":
        return _extract_pdf(file_bytes)
    elif ext == ".docx":
        return _extract_docx(file_bytes)
    else:
        raise ValueError(f"Unsupported file type '{ext}'. Only .pdf and .docx are supported.")


def _extract_pdf(file_bytes: bytes) -> str:
    import pdfplumber
    text_parts = []
    with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
        for page in pdf.pages:
            page_text = page.extract_text() or ""
            text_parts.append(page_text)
    text = "\n".join(text_parts).strip()
    if not text:
        raise ValueError(
            "No extractable text found in PDF. It may be a scanned/image-based PDF "
            "(would require OCR, which is out of scope for this tool)."
        )
    return text


def _extract_docx(file_bytes: bytes) -> str:
    import docx
    doc = docx.Document(io.BytesIO(file_bytes))
    parts = []
    for para in doc.paragraphs:
        if para.text.strip():
            parts.append(para.text)
    # Some resumes use tables for layout (e.g. skills grids) — capture those too.
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    parts.append(cell.text)
    text = "\n".join(parts).strip()
    if not text:
        raise ValueError("No extractable text found in DOCX file.")
    return text
